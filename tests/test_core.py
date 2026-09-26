"""Testes do núcleo. Rode com:  python -m pytest tests  (ou python tests/test_core.py)"""
import os
import sys
import tempfile

import pymupdf

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from gavetapdf.core import compress, convert, ocr, organize, sign  # noqa: E402
from gavetapdf.core.common import PdfError, parse_ranges  # noqa: E402
from gavetapdf import i18n  # noqa: E402

i18n.set_language("pt")  # os testes conferem mensagens e nomes de arquivo em português


def make_pdf(path, pages, text="Página"):
    doc = pymupdf.open()
    for i in range(pages):
        p = doc.new_page()
        p.insert_text((72, 100), f"{text} {i + 1}", fontsize=24)
    doc.save(path)
    doc.close()


def make_photo_pdf(path, pages=3):
    """PDF com imagens grandes (ruído colorido) para testar compressão."""
    doc = pymupdf.open()
    for _ in range(pages):
        pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 2400, 3300), False)
        pix.set_rect(pix.irect, (200, 180, 160))
        data = bytearray(pix.samples)
        import random
        rnd = random.Random(1)
        for k in range(0, len(data), 97):
            data[k] = rnd.randrange(256)
        pix = pymupdf.Pixmap(pymupdf.csRGB, 2400, 3300, bytes(data), False)
        page = doc.new_page()
        page.insert_image(page.rect, stream=pix.tobytes("png"))
    doc.save(path)
    doc.close()


def make_scanned_pdf(path, text):
    """Simula um documento escaneado: texto renderizado como imagem, sem camada de texto."""
    src = pymupdf.open()
    p = src.new_page()
    p.insert_textbox(pymupdf.Rect(72, 72, 520, 400), text, fontsize=16, fontname="helv")
    pix = p.get_pixmap(dpi=200)
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_image(page.rect, stream=pix.tobytes("png"))
    doc.save(path)


def test_parse_ranges():
    assert parse_ranges("1-3, 5, 8-", 10) == [[0, 1, 2], [4], [7, 8, 9]]
    assert parse_ranges("-2", 5) == [[0, 1]]
    for bad in ["0", "4-2", "11", "abc", ""]:
        try:
            parse_ranges(bad, 10)
        except PdfError:
            continue
        raise AssertionError(bad)


def test_merge_split_extract_build():
    with tempfile.TemporaryDirectory() as d:
        a, b = os.path.join(d, "a.pdf"), os.path.join(d, "b.pdf")
        make_pdf(a, 3, "A")
        make_pdf(b, 2, "B")
        merged = organize.merge([a, b], os.path.join(d, "m.pdf"))
        with pymupdf.open(merged) as m:
            assert m.page_count == 5
            assert "B 1" in m[3].get_text()

        parts = organize.split_ranges(merged, "1-2, 3, 4-", os.path.join(d, "out"))
        assert [pymupdf.open(p).page_count for p in parts] == [2, 1, 2]

        each = organize.split_every(merged, 1, os.path.join(d, "each"))
        assert len(each) == 5

        ex = organize.extract(merged, "5, 1", os.path.join(d, "ex.pdf"))
        with pymupdf.open(ex) as e:
            assert "B 2" in e[0].get_text() and "A 1" in e[1].get_text()

        refs = [organize.PageRef(b, 1, 90), organize.PageRef(a, 0, 0)]
        built = organize.build_from_pages(refs, a)  # salva por cima de uma das origens
        with pymupdf.open(built) as bt:
            assert bt.page_count == 2 and bt[0].rotation == 90
            assert "B 2" in bt[0].get_text()


def test_compress():
    with tempfile.TemporaryDirectory() as d:
        src = os.path.join(d, "fotos.pdf")
        make_photo_pdf(src)
        r = compress.compress(src, os.path.join(d, "c.pdf"), "forte")
        print(f"compressão: {r.before} -> {r.after} ({r.saved_percent:.0f}%)")
        assert r.after < r.before
        with pymupdf.open(r.out) as c:
            assert c.page_count == 3


def test_convert():
    with tempfile.TemporaryDirectory() as d:
        src = os.path.join(d, "doc.pdf")
        make_pdf(src, 3)
        imgs = convert.pdf_to_images(src, os.path.join(d, "img"), "jpg", 100, "1-2")
        assert len(imgs) == 2 and all(os.path.getsize(i) > 0 for i in imgs)
        pdf = convert.images_to_pdf(imgs, os.path.join(d, "back.pdf"), "a4", 10)
        with pymupdf.open(pdf) as p:
            assert p.page_count == 2
            assert abs(p[0].rect.width - 595) < 1
        txt = convert.pdf_to_text(src, os.path.join(d, "doc.txt"))
        assert "Página 3" in open(txt, encoding="utf-8-sig").read()


def test_ocr():
    if "por" not in ocr.available_languages():
        print("OCR: idioma 'por' ausente, teste ignorado")
        return
    with tempfile.TemporaryDirectory() as d:
        src = os.path.join(d, "scan.pdf")
        make_scanned_pdf(src, "Prefeitura Municipal\nCertidão de regularização\nAtenção ao prazo de emissão")
        with pymupdf.open(src) as s:
            assert s[0].get_text().strip() == ""
        out = os.path.join(d, "scan-ocr.pdf")
        stats = ocr.ocr_file(src, out, ["por"])
        with pymupdf.open(out) as o:
            text = o[0].get_text()
            print("OCR:", text.strip().replace("\n", " | "))
            assert "Prefeitura" in text and "Certidão" in text
            assert abs(o[0].rect.width - 595) < 2
        assert stats["ocr_pages"] == 1


def test_sign():
    with tempfile.TemporaryDirectory() as d:
        src = os.path.join(d, "doc.pdf")
        make_pdf(src, 3)
        with pymupdf.open(src) as doc:
            doc[1].set_rotation(90)
            doc.saveIncr()

        # "foto" da assinatura: traço escuro sobre papel levemente acinzentado
        tmp = pymupdf.open()
        p = tmp.new_page(width=400, height=200)
        p.draw_rect(p.rect, color=None, fill=(0.94, 0.94, 0.92))
        p.draw_bezier((40, 150), (120, 20), (220, 190), (360, 60), color=(0.1, 0.1, 0.4), width=5)
        photo = p.get_pixmap(dpi=144).tobytes("jpg")
        png = sign.clean_signature_image(photo)
        sig = pymupdf.Pixmap(png)
        assert sig.alpha and sig.pixel(0, 0)[3] == 0  # fundo transparente
        assert sig.width < 800  # bordas vazias recortadas

        out = os.path.join(d, "assinado.pdf")
        sign.sign(src, png, [sign.Placement(0, (300, 650, 500, 720)),
                             sign.Placement(1, (500, 400, 700, 470))], out)
        with pymupdf.open(out) as doc:
            assert [len(pg.get_images()) for pg in doc] == [1, 1, 0]
            # na página girada, a imagem fica onde o usuário a viu na tela
            bbox = doc[1].get_image_rects(doc[1].get_images()[0][0])[0] * doc[1].rotation_matrix
            assert abs(bbox.x0 - 500) < 1 and abs(bbox.y0 - 400) < 1

        for bad in ([], [sign.Placement(9, (0, 0, 10, 10))]):
            try:
                sign.sign(src, png, bad, out)
            except PdfError:
                continue
            raise AssertionError("deveria ter falhado")
        blank = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 60, 30), False)
        blank.set_rect(blank.irect, (255, 255, 255))
        try:
            sign.clean_signature_image(blank.tobytes("png"))
        except PdfError:
            pass
        else:
            raise AssertionError("imagem em branco deveria falhar")


def make_table_pdf(path):
    """Página com título, parágrafo e uma tabela com bordas (valores em reais)."""
    doc = pymupdf.open()
    p = doc.new_page()
    p.insert_text((72, 72), "Relatório de despesas", fontsize=18)
    p.insert_textbox(pymupdf.Rect(72, 90, 520, 150), "Despesas do trimestre em reais. " * 4, fontsize=10)
    rows = [["Item", "Qtd", "Valor"], ["Papel A4", "10", "R$ 1.250,00"], ["Toner", "2", "R$ 480,50"],
            ["Código", "0012", "15%"]]
    for i, row in enumerate(rows):
        for j, text in enumerate(row):
            rect = pymupdf.Rect(72 + j * 150, 180 + i * 22, 222 + j * 150, 202 + i * 22)
            p.draw_rect(rect, color=(0, 0, 0), width=0.8)
            p.insert_text((rect.x0 + 4, rect.y1 - 7), text, fontsize=10)
    doc.save(path)
    doc.close()


def make_web_table_pdf(path):
    """Tabela de página da web impressa: cada linha tem 4 células preenchidas com larguras
    próprias (diferentes do cabeçalho), células vazias e uma linha de largura total."""
    doc = pymupdf.open()
    p = doc.new_page()
    rows = [
        ([80, 120, 100, 200], ["Órgão", "Valor", "Data", "Objeto"]),
        ([60, 90, 150, 200], ["Prefeitura", "1.500,00", "", "Compra de papel"]),
        (None, ["Sem propostas"]),
        ([100, 70, 110, 220], ["Câmara", "", "01-02-2026", "Serviço de limpeza"]),
    ]
    y = 50
    for widths, texts in rows:
        widths = widths or [500]
        x = 50
        for w, text in zip(widths, texts + [""] * (len(widths) - len(texts))):
            p.draw_rect(pymupdf.Rect(x, y, x + w, y + 30), color=None, fill=(0.95, 0.95, 0.95))
            if text:
                p.insert_text((x + 4, y + 19), text, fontsize=9)
            x += w
        y += 30
    doc.save(path)
    doc.close()


def test_pdf_to_excel_web_table():
    from openpyxl import load_workbook

    with tempfile.TemporaryDirectory() as d:
        src = os.path.join(d, "web.pdf")
        make_web_table_pdf(src)
        xlsx = os.path.join(d, "web.xlsx")
        assert convert.pdf_to_excel(src, xlsx) == 1
        rows = [[c.value for c in row] for row in load_workbook(xlsx).active.iter_rows()]
        assert rows == [
            ["Órgão", "Valor", "Data", "Objeto"],
            ["Prefeitura", 1500.0, None, "Compra de papel"],
            ["Sem propostas", None, None, None],
            ["Câmara", None, "01-02-2026", "Serviço de limpeza"],
        ]

        # PDF → Texto: cada linha da tabela vira um registro "Coluna: valor", sem misturar colunas
        txt = os.path.join(d, "web.txt")
        convert.pdf_to_text(src, txt)
        text = open(txt, encoding="utf-8-sig").read()
        assert "Órgão:  Prefeitura\nValor:  1.500,00\nObjeto: Compra de papel" in text
        assert "\n\nSem propostas\n\n" in text
        assert "Órgão:  Câmara\nData:   01-02-2026\nObjeto: Serviço de limpeza" in text


def test_to_number():
    assert convert.to_number("1.250,00") == (1250.0, "#,##0.00")
    assert convert.to_number("R$ 480,50") == (480.5, '"R$" #,##0.00')
    assert convert.to_number("15%") == (0.15, "0.00%")
    assert convert.to_number("(1.000,00)")[0] == -1000.0
    assert convert.to_number("10") == (10, "General")
    for text in ["0012", "123.456.789-00", "01/02/2026", "abc", "", "7891234567890123", "-"]:
        assert convert.to_number(text) is None, text
    # formatos americano e russo
    assert convert.to_number("$1,234.56", ".") == (1234.56, '"$" #,##0.00')
    assert convert.to_number("1,234", ".") == (1234, "#,##0")
    assert convert.to_number("1 234,56 ₽") == (1234.56, '#,##0.00 "₽"')
    assert convert.guess_decimal(["1,234.50", "3.20", "Item"]) == "."
    assert convert.guess_decimal(["1.234,50", "3,20"], ".") == ","


def test_pdf_to_excel_and_word():
    from openpyxl import load_workbook
    import docx

    with tempfile.TemporaryDirectory() as d:
        src = os.path.join(d, "tabela.pdf")
        make_table_pdf(src)

        xlsx = os.path.join(d, "tabela.xlsx")
        assert convert.pdf_to_excel(src, xlsx) == 1
        ws = load_workbook(xlsx).active
        assert ws.title == "Página 1"
        assert [c.value for c in ws[1]] == ["Item", "Qtd", "Valor"]
        assert ws["C2"].value == 1250.0 and ws["B2"].value == 10
        assert ws["B4"].value == "0012" and ws["C4"].value == 0.15

        out = os.path.join(d, "tabela.docx")
        convert.pdf_to_word(src, out)
        document = docx.Document(out)
        assert "Relatório de despesas" in " ".join(p.text for p in document.paragraphs)
        assert document.tables and document.tables[0].rows[1].cells[0].text == "Papel A4"

        # PDF só com imagem (escaneado): mensagem pedindo OCR
        scanned = os.path.join(d, "scan.pdf")
        make_scanned_pdf(scanned, "Documento escaneado")
        for fn, target in ((convert.pdf_to_word, out), (convert.pdf_to_excel, xlsx)):
            try:
                fn(scanned, target)
            except PdfError as exc:
                assert "OCR" in str(exc)
            else:
                raise AssertionError("deveria pedir OCR")

        # texto sem tabelas
        plain = os.path.join(d, "texto.pdf")
        make_pdf(plain, 1)
        try:
            convert.pdf_to_excel(plain, xlsx)
        except PdfError as exc:
            assert "Nenhuma tabela" in str(exc)
        else:
            raise AssertionError("deveria avisar que não há tabelas")


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_"):
            fn()
            print("ok", name)
