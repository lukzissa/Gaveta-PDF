"""PDF escaneado com OCR → Word com texto de verdade (sem as fotos das páginas).

Num PDF "pesquisável" cada página é a foto do papel com o texto do OCR invisível por
cima. Converter isso como um PDF comum leva só as fotos para o Word. Aqui o documento
é remontado a partir das palavras do OCR, como se tivesse sido digitado no Word:
linhas viram parágrafos corridos, com alinhamento, recuos, tamanho de letra, negrito
(medido pela espessura da tinta na imagem) e fundo cinza de títulos (medido pela cor
atrás da linha).
"""
from __future__ import annotations

import re
import statistics
from dataclasses import dataclass, field

import pymupdf

from .common import NULL_CONTEXT, Context

INK = 128  # pixel mais escuro que isso é tinta
RENDER_DPI = 110


# ------------------------------------------------------------------ detecção
def is_ocr_page(page: pymupdf.Page) -> bool:
    """Página escaneada com texto de OCR: uma imagem cobre a página e o texto é invisível."""
    area = abs(page.rect)
    if not area or not any(abs(pymupdf.Rect(i["bbox"]) & page.rect) >= area * 0.8 for i in page.get_image_info()):
        return False
    invisible = visible = 0
    for span in page.get_texttrace():
        n = len(span.get("chars", ()))
        if span.get("type") == 3 or span.get("opacity", 1) == 0:
            invisible += n
        else:
            visible += n
    return invisible > 0 and invisible >= visible * 4


# ------------------------------------------------------------------ estrutura
@dataclass
class Word:
    rect: pymupdf.Rect
    text: str
    bold: bool = False


@dataclass
class Line:
    words: list[Word]
    shade: float | None = None  # tom de cinza do fundo (0-255), se houver faixa

    @property
    def x0(self): return min(w.rect.x0 for w in self.words)
    @property
    def x1(self): return max(w.rect.x1 for w in self.words)
    @property
    def y0(self): return statistics.median(w.rect.y0 for w in self.words)
    @property
    def y1(self): return statistics.median(w.rect.y1 for w in self.words)
    @property
    def text(self): return " ".join(w.text for w in self.words)


@dataclass
class Paragraph:
    lines: list[Line] = field(default_factory=list)
    align: str = "left"  # left, center, right, justify
    left_indent: float = 0.0
    first_indent: float = 0.0
    size: float = 11.0
    word_height: float = 0.0
    pitch: float | None = None  # distância entre linhas (parágrafos com mais de uma linha)
    space_after: float = 0.0
    shade: float | None = None


def _ink_and_background(pix: pymupdf.Pixmap, rect: pymupdf.Rect, scale: float) -> tuple[float, float]:
    """Proporção de pixels de tinta e tom mediano do fundo dentro de `rect` (coordenadas da página)."""
    x0, y0 = max(int(rect.x0 * scale), 0), max(int(rect.y0 * scale), 0)
    x1, y1 = min(int(rect.x1 * scale) + 1, pix.width), min(int(rect.y1 * scale) + 1, pix.height)
    if x1 <= x0 or y1 <= y0:
        return 0.0, 255.0
    data, stride = pix.samples, pix.stride
    ink, light = 0, []
    for y in range(y0, y1):
        row = data[y * stride + x0:y * stride + x1]
        for v in row:
            if v < INK:
                ink += 1
            else:
                light.append(v)
    total = (x1 - x0) * (y1 - y0)
    return ink / total, (statistics.median(light) if light else 255.0)


_JUNK = set("|[]{}~_\\^`'\"“”‘’¦")


def _is_junk_word(text: str) -> bool:
    """Símbolos soltos que o OCR "lê" em rabiscos, bordas e rubricas."""
    t = text.strip()
    return not t or all(c in _JUNK for c in t) or (not any(c.isalnum() for c in t) and len(t) > 2)


def _word_height(words: list[Word]) -> float:
    """Altura típica das palavras (ignora caixas do OCR com altura absurda)."""
    hs = sorted(w.rect.height for w in words if sum(c.isalnum() for c in w.text) >= 2) or \
        sorted(w.rect.height for w in words)
    return hs[len(hs) // 3]  # um pouco abaixo da mediana: caixas erradas são sempre maiores


def _page_lines(page: pymupdf.Page) -> list[Line]:
    raw = [w for w in page.get_text("words") if not _is_junk_word(w[4])]
    if not raw:
        return []
    # linhas como o OCR as reconheceu (bloco, linha): as caixas de algumas palavras vêm
    # com altura errada, mas a qual linha cada palavra pertence ele acerta
    grouped: dict[tuple[int, int], list] = {}
    for w in raw:
        grouped.setdefault((w[5], w[6]), []).append(Word(pymupdf.Rect(w[:4]), w[4]))
    lines = [Line(sorted(ws, key=lambda w: w.rect.x0)) for ws in grouped.values()]
    h = _word_height([w for ln in lines for w in ln.words])

    # texto justificado com espaços largos às vezes vira várias "linhas" na mesma altura: junta
    lines.sort(key=lambda ln: ((ln.y0 + ln.y1) / 2, ln.x0))
    merged: list[Line] = []
    for ln in lines:
        cy = (ln.y0 + ln.y1) / 2
        for other in reversed(merged[-4:]):
            same_row = abs((other.y0 + other.y1) / 2 - cy) < h * 0.65
            apart = ln.x0 >= other.x1 - 2 or ln.x1 <= other.x0 + 2
            if same_row and apart:
                other.words = sorted(other.words + ln.words, key=lambda w: w.rect.x0)
                break
        else:
            merged.append(ln)
    lines = sorted(merged, key=lambda ln: (ln.y0, ln.x0))

    scale = RENDER_DPI / 72
    pix = page.get_pixmap(dpi=RENDER_DPI, colorspace=pymupdf.csGRAY, alpha=False)
    ink: dict[int, float] = {}
    for ln in lines:
        top, bottom = ln.y0, ln.y1
        for w in ln.words:
            # mede a tinta só na faixa da linha (a caixa da palavra pode vir errada)
            ink[id(w)], _ = _ink_and_background(pix, pymupdf.Rect(w.rect.x0, top, w.rect.x1, bottom), scale)
        _, background = _ink_and_background(pix, pymupdf.Rect(ln.x0, top, ln.x1, bottom), scale)
        ln.shade = background if background < 235 else None

    # negrito: tinta bem acima do normal do documento, depois suavizado por linha
    letters = lambda w: sum(c.isalpha() for c in w.text)  # noqa: E731
    densities = [ink[id(w)] for ln in lines for w in ln.words if letters(w) >= 3]
    normal = statistics.median(densities) if densities else 0
    for ln in lines:
        for w in ln.words:
            w.bold = bool(normal) and ink[id(w)] > normal * 1.28 and letters(w) > 0
        long_words = [w for w in ln.words if letters(w) >= 3]
        share = sum(w.bold for w in long_words) / len(long_words) if long_words else 0
        if share >= 0.6:
            for w in ln.words:
                w.bold = True
        elif share <= 0.15:
            for w in ln.words:
                w.bold = False
        else:  # palavras curtas ("e", "de", "nº") seguem as vizinhas
            for i, w in enumerate(ln.words):
                if letters(w) < 3:
                    before = ln.words[i - 1].bold if i else False
                    after = ln.words[i + 1].bold if i + 1 < len(ln.words) else False
                    w.bold = before and after
    return [ln for ln in lines if not _is_noise(ln, page)]


def _is_noise(line: Line, page: pymupdf.Page) -> bool:
    text = line.text.strip()
    alnum = sum(c.isalnum() for c in text)
    if alnum == 0:
        return True
    # número de página sozinho no rodapé
    if line.y0 > page.rect.height * 0.9 and len(text) <= 6 and alnum <= 3:
        return True
    return False


_LIST_ITEM = re.compile(r"^(\(?[a-zA-Z0-9]{1,3}[\)\.\-–]|[•\-–])\s")


def _paragraphs(lines: list[Line], page: pymupdf.Page) -> tuple[list[Paragraph], float, float]:
    if not lines:
        return [], 0, page.rect.width
    wide = [ln for ln in lines if ln.x1 - ln.x0 > page.rect.width * 0.35] or lines
    left = sorted(ln.x0 for ln in wide)[len(wide) // 10]
    right = sorted(ln.x1 for ln in wide)[-(len(wide) // 10) - 1]
    width = max(right - left, 1)
    heights = [ln.y1 - ln.y0 for ln in lines]
    h = statistics.median(heights)

    # rabiscos e rubricas nas margens não são texto do documento: palavras fora da área do
    # texto saem, e linhas curtas soltas na metade direita (rubricas no canto) também
    for ln in lines:
        ln.words = [w for w in ln.words if left - 4 <= w.rect.x0 and w.rect.x1 <= right + 4] or ln.words
        while len(ln.words) > 1:
            last, before = ln.words[-1], ln.words[-2]
            if sum(c.isalnum() for c in last.text) <= 2 and last.rect.x0 - before.rect.x1 > h * 1.5:
                ln.words.pop()  # letra solta longe do texto: rubrica na margem
            else:
                break
    lines = [ln for ln in lines if not (
        len(ln.text) < 25 and (ln.x0 > right - 3 or ln.x1 < left + 3 or
                               (ln.x0 > left + width * 0.55 and sum(c.isalnum() for c in ln.text) <= 6)))]
    if not lines:
        return [], left, right

    def kind(ln: Line) -> str:
        gap_left, gap_right = ln.x0 - left, right - ln.x1
        if min(gap_left, gap_right) > width * 0.05 and abs(gap_left - gap_right) < width * 0.045:
            return "center"
        if gap_right < width * 0.03 and gap_left > width * 0.3:
            return "right"
        return "left"

    paras: list[Paragraph] = []
    prev = None
    for ln in lines:
        start = prev is None
        if prev is not None:
            gap = ln.y0 - prev.y1
            short_prev = prev.x1 < right - width * 0.12
            same_block = kind(ln) == kind(prev) or (kind(prev) == "right" and kind(ln) != "center")
            start = (gap > h * 0.9 or short_prev or not same_block
                     or _LIST_ITEM.match(ln.text) is not None
                     or (ln.shade is not None) != (prev.shade is not None))
        if start:
            paras.append(Paragraph(shade=ln.shade))
        paras[-1].lines.append(ln)
        prev = ln

    for i, p in enumerate(paras):
        kinds = {kind(ln) for ln in p.lines}
        first, rest = p.lines[0], p.lines[1:]
        if kinds == {"center"}:
            p.align = "center"
        elif kinds == {"right"} and not rest:
            p.align = "right"
        else:  # inclui blocos recuados para a direita com várias linhas (justificados)
            p.align = "justify" if rest and sum(ln.x1 > right - width * 0.04 for ln in p.lines[:-1]) >= len(p.lines[:-1]) * 0.6 else "left"
            base = min(ln.x0 for ln in rest) if rest else first.x0
            p.left_indent = max(base - left, 0) if base - left > width * 0.02 else 0
            p.first_indent = first.x0 - (left + p.left_indent)
            if abs(p.first_indent) < width * 0.015:
                p.first_indent = 0
        pitch = [b.y0 - a.y0 for a, b in zip(p.lines, p.lines[1:]) if 0 < b.y0 - a.y0 < h * 3]
        p.pitch = statistics.median(pitch) if pitch else None
        p.word_height = _word_height([w for ln in p.lines for w in ln.words])
        if i + 1 < len(paras):
            p.space_after = max(0.0, min(paras[i + 1].lines[0].y0 - p.lines[-1].y1 - h * 0.35, 24.0))
    return paras, left, right


def _join(lines: list[Line]) -> list[Word]:
    """Palavras do parágrafo em ordem, desfazendo hifenização de fim de linha."""
    out: list[Word] = []
    for ln in lines:
        for w in ln.words:
            if out and out[-1].text.endswith("-") and w is ln.words[0] and w.text[:1].islower():
                out[-1] = Word(out[-1].rect, out[-1].text[:-1] + w.text, out[-1].bold or w.bold)
            else:
                out.append(Word(w.rect, w.text, w.bold))
    return out


def _font_metrics(name: str) -> pymupdf.Font:
    """Medidas da fonte usada no Word (arquivo do Windows); Helvetica se não existir."""
    import os

    files = {"Calibri": "calibri.ttf", "Arial": "arial.ttf"}
    path = os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts", files.get(name, ""))
    return pymupdf.Font(fontfile=path) if name in files and os.path.exists(path) else pymupdf.Font("helv")


def text_only_copy(doc: pymupdf.Document, scanned: list[bool]) -> pymupdf.Document:
    """Cópia do PDF em que as páginas escaneadas trazem só o texto do OCR, visível e sem a foto."""
    import os

    arial = os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts", "arial.ttf")
    font_args = {"fontfile": arial, "fontname": "Arial"} if os.path.exists(arial) else {"fontname": "helv"}
    font = pymupdf.Font(fontfile=arial) if os.path.exists(arial) else pymupdf.Font("helv")
    result = pymupdf.open()
    for page, is_scan in zip(doc, scanned):
        if not is_scan:
            result.insert_pdf(doc, from_page=page.number, to_page=page.number)
            continue
        new = result.new_page(width=page.rect.width, height=page.rect.height)
        for ln in _page_lines(page):
            size = max(6.0, min((ln.y1 - ln.y0) * 0.85, 30.0))
            for w in ln.words:
                width = font.text_length(w.text, size)
                fs = size * min(1.0, w.rect.width / width) if width else size
                new.insert_text((w.rect.x0, ln.y1 - (ln.y1 - ln.y0) * 0.2), w.text, fontsize=fs, **font_args)
    return result


# ------------------------------------------------------------------ Word
def ocr_pdf_to_docx(doc: pymupdf.Document, out: str, ctx: Context = NULL_CONTEXT, font: str = "Calibri") -> None:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    from ..i18n import tr

    pages = []
    for i, page in enumerate(doc):
        ctx.progress(i, doc.page_count + 1, tr("Página {0} de {1}", i + 1, doc.page_count))
        lines = _page_lines(page)
        paras, left, right = _paragraphs(lines, page)
        pages.append((page.rect, paras, left, right, lines))

    # Tamanho da letra único: para cada linha que NÃO é justificada (termina antes da
    # margem direita), calcula que tamanho a fonte escolhida precisa ter para a linha ocupar
    # a mesma largura que no papel; vale a mediana. As caixas das palavras do OCR e o
    # espaçamento entre linhas não servem para isso, e em documentos escaneados quase tudo
    # tem o mesmo tamanho.
    all_paras = [p for _, ps, *_ in pages for p in ps]
    metrics = _font_metrics(font)
    fits = []
    for _, ps, left, right, _ in pages:
        for p in ps:
            for ln in p.lines:
                em = metrics.text_length(ln.text, 1)
                if len(ln.text) >= 25 and em and ln.x1 < right - (right - left) * 0.08:
                    fits.append((ln.x1 - ln.x0) / em)
    if fits:
        body = statistics.median(fits)
    else:
        pitches = [p.pitch for p in all_paras if p.pitch]
        body = statistics.median(pitches) / 1.35 if pitches else 11.0
    body = max(8.0, min(round(body * 2) / 2, 14.0))
    for p in all_paras:
        p.size = body

    out_doc = docx.Document()
    style = out_doc.styles["Normal"]
    style.font.name = font
    style.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    style.font.size = Pt(body)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.space_before = Pt(0)

    first_rect, _, first_left, first_right, _ = pages[0]
    tops = [ls[0].y0 for *_, ls in pages if ls]
    bottoms = [first_rect.height - ls[-1].y1 for *_, ls in pages if ls]
    sec = out_doc.sections[0]
    sec.page_width, sec.page_height = Pt(first_rect.width), Pt(first_rect.height)
    sec.left_margin = Pt(max(28, min(first_left, 110)))
    sec.right_margin = Pt(max(28, min(first_rect.width - first_right, 110)))
    sec.top_margin = Pt(max(28, min(min(tops) if tops else 72, 100)))
    sec.bottom_margin = Pt(max(28, min(min(bottoms) if bottoms else 72, 100)))

    aligns = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
              "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}

    # parágrafo que continua na página seguinte vira um só
    flat: list[Paragraph] = []
    for _, ps, *_ in pages:
        for j, p in enumerate(ps):
            if j == 0 and flat:
                last = flat[-1]
                tail = last.lines[-1].text.rstrip()
                head = p.lines[0].text.lstrip()
                if tail and tail[-1] not in ".:;!?)" and head[:1].islower() and last.align == p.align:
                    last.lines.extend(p.lines)
                    last.space_after = p.space_after
                    continue
            flat.append(p)

    # recuos parecidos (o OCR varia alguns pontos de uma linha para outra) viram o mesmo valor
    for attr in ("left_indent", "first_indent"):
        values = sorted({round(getattr(p, attr)) for p in flat if getattr(p, attr)})
        groups: list[list[int]] = []
        for v in values:
            if groups and v - groups[-1][-1] <= 6:
                groups[-1].append(v)
            else:
                groups.append([v])
        snap = {v: statistics.median(g) for g in groups for v in g}
        for p in flat:
            if getattr(p, attr):
                setattr(p, attr, snap[round(getattr(p, attr))])
    gaps = [p.space_after for p in flat if p.space_after > 2]
    usual_gap = statistics.median(gaps) if gaps else 8.0
    for a, b in zip(flat, flat[1:]):
        if a.space_after < 1 and b.lines[0].y0 < a.lines[-1].y0:  # b começa numa página nova
            a.space_after = usual_gap
        if b.shade is not None:  # título com faixa: respiro antes
            a.space_after = max(a.space_after, usual_gap)

    for p in flat:
        para = out_doc.add_paragraph()
        fmt = para.paragraph_format
        para.alignment = aligns[p.align]
        if p.left_indent:
            fmt.left_indent = Pt(p.left_indent)
        if p.first_indent:
            fmt.first_line_indent = Pt(p.first_indent)
        fmt.space_after = Pt(round(p.space_after))
        if p.shade is not None:
            shd = OxmlElement("w:shd")
            tone = f"{int(p.shade):02X}" * 3
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), tone)
            fmt._element.get_or_add_pPr().append(shd)
        words = _join(p.lines)
        long_words = [w for w in words if sum(c.isalpha() for c in w.text) >= 3]
        if long_words and sum(w.bold for w in long_words) >= len(long_words) * 0.7:
            for w in words:
                w.bold = True
        # junta palavras vizinhas com o mesmo estilo num só trecho de texto
        runs: list[tuple[bool, list[str]]] = []
        for w in words:
            if runs and runs[-1][0] == w.bold:
                runs[-1][1].append(w.text)
            else:
                runs.append((w.bold, [w.text]))
        for k, (bold, texts) in enumerate(runs):
            run = para.add_run((" " if k else "") + " ".join(texts))
            run.bold = bold or None
            if p.size != body:
                run.font.size = Pt(p.size)
    if not flat:
        out_doc.add_paragraph()
    out_doc.save(out)
