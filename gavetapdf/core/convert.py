"""Conversões: PDF → imagens/texto/Word/Excel e imagens → PDF."""
from __future__ import annotations

import html
import logging
import os
import re
import textwrap
import warnings

import pymupdf

from .common import NULL_CONTEXT, Context, PdfError, open_pdf, parse_ranges, save_pdf, stem, ticking, unique_path
from ..i18n import decimal_separator, tr

# Tamanho A4 em pontos (1/72 pol.)
A4 = pymupdf.paper_rect("a4")


def pdf_to_images(
    path: str,
    out_dir: str,
    fmt: str = "png",
    dpi: int = 150,
    ranges_text: str = "",
    ctx: Context = NULL_CONTEXT,
) -> list[str]:
    fmt = fmt.lower()
    if fmt not in ("png", "jpg"):
        raise PdfError(tr("Formato de imagem não suportado."))
    outputs: list[str] = []
    os.makedirs(out_dir, exist_ok=True)
    with open_pdf(path) as doc:
        if ranges_text.strip():
            pages = [p for g in parse_ranges(ranges_text, doc.page_count) for p in g]
        else:
            pages = list(range(doc.page_count))
        width = len(str(doc.page_count))
        for i, p in enumerate(pages):
            ctx.progress(i, len(pages), tr("Página {0}", p + 1))
            pix = doc[p].get_pixmap(dpi=dpi, alpha=False)
            name = tr("{0} - página {1}.{2}", stem(path), str(p + 1).zfill(width), fmt)
            out = unique_path(os.path.join(out_dir, name))
            if fmt == "jpg":
                pix.save(out, jpg_quality=90)
            else:
                pix.save(out)
            outputs.append(out)
    ctx.progress(1, 1, tr("Concluído"))
    return outputs


def images_to_pdf(
    images: list[str],
    out: str,
    page_size: str = "original",
    margin_mm: float = 0,
    ctx: Context = NULL_CONTEXT,
) -> str:
    """Cria um PDF com uma imagem por página.

    page_size: "original" (página do tamanho da imagem) ou "a4" (ajusta e centraliza,
    virando a folha para paisagem quando a imagem é mais larga que alta).
    """
    if not images:
        raise PdfError(tr("Adicione ao menos uma imagem."))
    result = pymupdf.open()
    margin = margin_mm * 72 / 25.4
    for i, img_path in enumerate(images):
        ctx.progress(i, len(images), tr("Adicionando {0}", os.path.basename(img_path)))
        try:
            img = pymupdf.open(img_path)
        except Exception as exc:  # noqa: BLE001
            raise PdfError(tr("Não foi possível ler a imagem:\n{0}", os.path.basename(img_path))) from exc
        with img:
            # GIF/TIFF podem ter vários quadros: usa todos.
            frames = img.page_count
            img_pdf = pymupdf.open("pdf", img.convert_to_pdf())
        with img_pdf:
            for f in range(frames):
                rect = img_pdf[f].rect
                if page_size == "a4":
                    page_rect = A4 if rect.height >= rect.width else pymupdf.Rect(0, 0, A4.height, A4.width)
                    page = result.new_page(width=page_rect.width, height=page_rect.height)
                    area = page_rect + (margin, margin, -margin, -margin)
                    page.show_pdf_page(area, img_pdf, f, keep_proportion=True)
                else:
                    page = result.new_page(width=rect.width + 2 * margin, height=rect.height + 2 * margin)
                    page.show_pdf_page(page.rect + (margin, margin, -margin, -margin), img_pdf, f)
    ctx.progress(len(images), len(images), tr("Salvando…"))
    save_pdf(result, out)
    result.close()
    return out


def pdf_to_text(path: str, out: str, ctx: Context = NULL_CONTEXT) -> str:
    chunks: list[str] = []
    with open_pdf(path) as doc:
        for i, page in enumerate(doc):
            ctx.progress(i, doc.page_count, tr("Página {0}", i + 1))
            chunks.append(tr("===== Página {0} =====\n{1}\n", i + 1, _page_as_text(page)))
    text = "\n".join(chunks)
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    with open(out, "w", encoding="utf-8-sig") as fh:  # BOM: o Bloco de Notas abre acentos corretamente
        fh.write(text)
    ctx.progress(1, 1, tr("Concluído"))
    return out


SCANNED_MESSAGE = (
    tr("Este PDF não tem texto selecionável (parece ser escaneado ou uma foto).\n\n"
    "Passe o arquivo pelo OCR antes e depois converta o resultado.")
)


def _has_text(doc: pymupdf.Document) -> bool:
    return any(page.get_text("text").strip() for page in doc)


def pdf_to_word(path: str, out: str, ctx: Context = NULL_CONTEXT) -> str:
    """Converte para .docx com texto editável.

    PDFs digitais vão pelo pdf2docx (parágrafos, tabelas e imagens). Páginas escaneadas
    com OCR (foto da página + texto invisível) são remontadas só com o texto reconhecido,
    senão o Word receberia apenas as fotos.
    """
    from . import ocr_docx

    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    temp = None
    with open_pdf(path) as doc:
        if not _has_text(doc):
            raise PdfError(SCANNED_MESSAGE)
        scanned = [ocr_docx.is_ocr_page(page) for page in doc]
        if all(s for s, page in zip(scanned, doc) if page.get_text("text").strip()):
            try:
                ocr_docx.ocr_pdf_to_docx(doc, out, ctx)
            except PermissionError as exc:
                raise PdfError(tr("Não foi possível salvar “{0}”.\nSe ele estiver aberto no Word, feche e tente de novo.", os.path.basename(out))) from exc
            ctx.progress(1, 1, tr("Concluído"))
            return out
        if any(scanned):
            # documento misto: as páginas escaneadas viram só texto antes de converter
            import tempfile

            fd, temp = tempfile.mkstemp(suffix=".pdf", dir=os.path.dirname(os.path.abspath(out)))
            os.close(fd)
            ocr_docx.text_only_copy(doc, scanned).save(temp)
    try:
        return _pdf2docx(temp or path, out, ctx, os.path.basename(path))
    finally:
        if temp:
            os.remove(temp)


def _pdf2docx(path: str, out: str, ctx: Context, name: str) -> str:
    """Conversão pelo pdf2docx (PDFs com texto digital). `name` é o nome mostrado nos erros."""
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        from pdf2docx import Converter

    # O pdf2docx registra cada passo no log; só interessam os erros.
    logging.disable(logging.INFO)
    cv = Converter(path)
    try:
        settings = cv.default_settings
        cv.load_pages()
        pages = list(cv.pages)
        # a barra é dividida entre a análise inicial (um passo só, ~0,13 s por página), as
        # páginas e a criação do .docx (outro passo só, ~0,06 s por página)
        n = len(pages)
        docx_part = max(n // 2, 1)
        total = n * 2 + docx_part
        with ticking(ctx, 0, n, total, 0.13 * n, tr("Analisando o documento…")):
            cv.parse_document(**settings)
        for i, page in enumerate(pages):
            ctx.progress(n + i, total, tr("Página {0} de {1}", i + 1, n))
            try:
                page.parse(**settings)
            except Exception as exc:  # noqa: BLE001 - uma página com problema não impede as outras
                logging.getLogger("gavetapdf").error("Página %d ignorada na conversão para Word: %s", i + 1, exc)
        os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
        try:
            with ticking(ctx, n * 2, docx_part, total, 0.06 * n, tr("Criando o documento Word…")):
                cv.make_docx(out, **settings)
        except PermissionError as exc:
            raise PdfError(tr("Não foi possível salvar “{0}”.\nSe ele estiver aberto no Word, feche e tente de novo.", os.path.basename(out))) from exc
        except Exception as exc:  # noqa: BLE001
            raise PdfError(tr("Não foi possível converter “{0}” para Word.\n\n{1}", name, exc)) from exc
    finally:
        cv.close()
        logging.disable(logging.NOTSET)
    ctx.progress(1, 1, tr("Concluído"))
    return out


# Vírgula decimal (Brasil, Espanha, Rússia): "1.234,56" ou "1 234,56"
_NUMBER_COMMA = re.compile(r"-?(\d{1,3}([. ]\d{3})+|\d+)(,\d+)?")
# Ponto decimal (EUA, Reino Unido): "1,234.56"
_NUMBER_DOT = re.compile(r"-?(\d{1,3}(,\d{3})+|\d+)(\.\d+)?")
_CURRENCIES = ("US$", "R$", "руб.", "руб", "$", "€", "£", "₽")


def guess_decimal(values: list[str], default: str = ",") -> str:
    """Descobre se o documento usa vírgula ou ponto como separador decimal."""
    comma = sum(bool(re.search(r"\d,\d{1,2}$|\d\.\d{3},\d", v)) for v in values)
    dot = sum(bool(re.search(r"\d\.\d{1,2}$|\d,\d{3}\.\d", v)) for v in values)
    if comma == dot:
        return default
    return "," if comma > dot else "."


def to_number(text: str, decimal: str = ",") -> tuple[float | int, str] | None:
    """Reconhece números como "1.234,56", "R$ 12,90", "1 234,56 ₽", "$1,234.56", "15%", "-3".

    `decimal` é o separador decimal do documento ("," ou ".").
    Retorna (valor, formato do Excel) ou None se o texto não for um número.
    Zero à esquerda ("0012"), CPF, CEP, datas e códigos longos continuam como texto.
    """
    s = " ".join(text.replace(" ", " ").replace(" ", " ").split())
    prefix = suffix = ""
    for sym in _CURRENCIES:
        if s.startswith(sym):
            prefix, s = sym, s[len(sym):].strip()
            break
        if s.endswith(sym):
            suffix, s = sym, s[:-len(sym)].strip()
            break
    percent = s.endswith("%")
    if percent:
        s = s[:-1].strip()
    if s.startswith("(") and s.endswith(")"):  # (1.234,56) = negativo em relatórios contábeis
        s = "-" + s[1:-1].strip()
    pattern = _NUMBER_COMMA if decimal == "," else _NUMBER_DOT
    if not s or not pattern.fullmatch(s):
        return None
    digits = s.lstrip("-")
    if len(digits) > 1 and digits[0] == "0" and not digits.startswith("0" + decimal):
        return None
    plain = s
    for sep in ((".", " ") if decimal == "," else (",",)):
        plain = plain.replace(sep, "")
    if decimal in plain:
        value: float | int = float(plain.replace(decimal, "."))
        fmt = "#,##0." + "0" * min(len(plain.split(decimal)[1]), 6)
    else:
        value = int(plain)
        if abs(value) >= 10**15:
            return None
        fmt = "#,##0" if plain != s else "General"  # tinha separador de milhar
    if prefix or suffix:
        fmt = f'"{prefix}" #,##0.00' if prefix else f'#,##0.00 "{suffix}"'
    if percent:
        value = value / 100
        fmt = "0.00%"
    return value, fmt


def _cell_bands(page: pymupdf.Page) -> list[list[pymupdf.Rect]]:
    """Faixas horizontais de retângulos lado a lado (as células de uma linha), de cima para baixo."""
    page_area = abs(page.rect)
    rects = []
    for d in page.get_drawings():
        items = d.get("items") or []
        if len(items) != 1 or items[0][0] not in ("re", "qu"):
            continue
        r = pymupdf.Rect(d["rect"])
        if r.width < 3 or r.height < 3 or abs(r) > page_area * 0.8:
            continue
        rects.append(r)
    # mesmo retângulo desenhado duas vezes (preenchimento + borda) conta uma vez só
    unique = {(round(r.x0), round(r.y0), round(r.x1), round(r.y1)): r for r in rects}
    bands: list[list[pymupdf.Rect]] = []
    for r in sorted(unique.values(), key=lambda r: (round(r.y0), r.x0)):
        if bands and abs(bands[-1][0].y0 - r.y0) <= 1.5 and abs(bands[-1][0].y1 - r.y1) <= 1.5:
            bands[-1].append(r)
        else:
            bands.append([r])
    valid = []
    for band in bands:
        band.sort(key=lambda r: r.x0)
        if all(abs(b.x0 - a.x1) <= 2 for a, b in zip(band, band[1:])):  # células encostadas
            valid.append(band)
    return valid


def _cell_grid_tables(page: pymupdf.Page) -> list[tuple[pymupdf.Rect, list[list[str]]]]:
    """Tabelas desenhadas célula por célula, comuns em páginas da web salvas como PDF.

    Nelas cada linha pode ter larguras de coluna diferentes das do cabeçalho, então a
    posição horizontal não diz a coluna: vale a ordem da célula dentro da linha
    (1ª célula = 1ª coluna, e assim por diante). Linhas de uma célula só, da largura
    da tabela (ex.: "Sem propostas"), entram com o texto na primeira coluna.
    """
    words = page.get_text("words", sort=True)

    def band_text(band: list[pymupdf.Rect]) -> list[str]:
        lines: list[dict[tuple, list[str]]] = [{} for _ in band]
        for w in words:
            cx, cy = (w[0] + w[2]) / 2, (w[1] + w[3]) / 2
            for i, cell in enumerate(band):
                if cell.x0 <= cx <= cell.x1 and cell.y0 <= cy <= cell.y1:
                    lines[i].setdefault((w[5], w[6]), []).append(w[4])
                    break
        return [" ".join(" ".join(ws) for ws in cell.values()) for cell in lines]

    tables, current, columns, left, right = [], [], 0, 0.0, 0.0

    def close():
        if sum(len(r) > 1 for r in current) >= 2:
            rows = []
            for band in current:
                values = band_text(band)
                rows.append(values if len(band) > 1 else values + [""] * (columns - 1))
            rows = [r for r in rows if any(r)]
            if rows:
                area = pymupdf.Rect(current[0][0].x0, current[0][0].y0, current[-1][-1].x1, current[-1][-1].y1)
                tables.append((area, rows))

    for band in _cell_bands(page):
        x0, x1 = band[0].x0, band[-1].x1
        same_width = current and abs(x0 - left) <= 3 and abs(x1 - right) <= 3
        if len(band) >= 2 and same_width and len(band) == columns:
            current.append(band)
        elif len(band) == 1 and same_width:
            current.append(band)
        elif len(band) >= 2:
            close()
            current, columns, left, right = [band], len(band), x0, x1
        else:
            close()
            current, columns = [], 0
    close()
    return tables


def _located_tables(page: pymupdf.Page, strategy: str) -> list[tuple[pymupdf.Rect, list[list[str]]]]:
    """Tabelas da página com a área que cada uma ocupa."""
    if strategy == "lines":
        grid = _cell_grid_tables(page)
        if grid:
            return grid
    try:
        found = page.find_tables(strategy=strategy)
    except Exception:  # noqa: BLE001
        return []
    tables = []
    for t in found.tables:
        rows = [["" if c is None else str(c).strip() for c in row] for row in t.extract()]
        rows = [r for r in rows if any(r)]
        if rows and (len(rows) > 1 or len(rows[0]) > 1):
            tables.append((pymupdf.Rect(t.bbox), rows))
    return tables


def _page_tables(page: pymupdf.Page, strategy: str) -> list[list[list[str]]]:
    return [rows for _, rows in _located_tables(page, strategy)]


def _clean_cell(value: str) -> str:
    """Quebras de linha do PDF viram espaço; "&nbsp;" e afins de páginas da web viram texto."""
    return " ".join(html.unescape(value).split())


def _table_as_text(rows: list[list[str]], width: int = 100) -> str:
    """Tabela em texto legível no Bloco de Notas.

    Com cabeçalho, cada linha vira um registro "Coluna: valor" (campos vazios são omitidos
    e textos longos quebram alinhados). Linhas de uma célula só (ex.: "Sem propostas")
    saem como estão. Sem cabeçalho reconhecível, as células são separadas por " | ".
    """
    rows = [[_clean_cell(c) for c in row] for row in rows]
    header = rows[0]
    has_header = len(rows) > 1 and all(header) and max(len(h) for h in header) <= 40
    if not has_header:
        return "\n".join(" | ".join(c for c in row if c) for row in rows)
    label = max(len(h) for h in header) + 2
    records = []
    for row in rows[1:]:
        filled = [(h, v) for h, v in zip(header, row) if v]
        if len(filled) == 1 and row[0]:
            records.append(row[0])
            continue
        lines = []
        for h, v in filled:
            wrapped = textwrap.wrap(v, max(width - label, 30)) or [""]
            lines.append(f"{h + ':':<{label}}{wrapped[0]}")
            lines += [" " * label + part for part in wrapped[1:]]
        records.append("\n".join(lines))
    return "\n\n".join(records)


def _page_as_text(page: pymupdf.Page) -> str:
    """Texto da página na ordem de leitura, com as tabelas organizadas por registro."""
    tables = _located_tables(page, "lines")
    parts: list[tuple[float, float, str]] = [(area.y0, area.x0, _table_as_text(rows)) for area, rows in tables]
    for x0, y0, x1, y1, text, _, kind in page.get_text("blocks", sort=True):
        if kind != 0 or not text.strip():
            continue  # imagens e blocos vazios
        center = pymupdf.Point((x0 + x1) / 2, (y0 + y1) / 2)
        if any(center in area for area, _ in tables):
            continue  # já saiu na tabela
        parts.append((y0, x0, text.strip()))
    parts.sort(key=lambda p: (round(p[0]), p[1]))
    return "\n\n".join(p[2] for p in parts)


def pdf_to_excel(path: str, out: str, one_sheet: bool = False, numbers: bool = True,
                 ctx: Context = NULL_CONTEXT) -> int:
    """Extrai as tabelas do PDF para uma planilha .xlsx e retorna quantas foram encontradas.

    one_sheet=False: uma aba por página; True: todas as tabelas numa aba só.
    numbers=True: "1.234,56", "R$ 10,00" e "15%" viram números de verdade no Excel.
    """
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    with open_pdf(path) as doc:
        if not _has_text(doc):
            raise PdfError(SCANNED_MESSAGE)
        total = doc.page_count + 1
        per_page = []
        for i, page in enumerate(doc):
            ctx.progress(i, total, tr("Procurando tabelas na página {0}", i + 1))
            per_page.append(_page_tables(page, "lines"))
        if not any(per_page):
            # Tabelas sem linhas desenhadas: tenta pelo alinhamento do texto.
            for i, page in enumerate(doc):
                ctx.progress(i, total, tr("Procurando tabelas sem bordas na página {0}", i + 1))
                per_page[i] = _page_tables(page, "text")

    count = sum(len(t) for t in per_page)
    if not count:
        raise PdfError(tr("Nenhuma tabela foi encontrada em “{0}”.\n\nPara levar só o texto, use PDF → Word ou PDF → Texto.", os.path.basename(path)))

    ctx.progress(total - 1, total, tr("Criando a planilha…"))
    cells = [v for tables in per_page for t in tables for row in t[1:] for v in row]
    decimal = guess_decimal(cells, decimal_separator())
    wb = Workbook()
    wb.remove(wb.active)
    bold = Font(bold=True)
    top = Alignment(vertical="top")
    wrap = Alignment(wrap_text=True, vertical="top")

    def write(ws, row: int, table: list[list[str]]) -> int:
        for r, values in enumerate(table):
            for c, value in enumerate(values):
                value = _clean_cell(value)
                cell = ws.cell(row=row + r, column=c + 1)
                cell.alignment = top
                parsed = to_number(value, decimal) if numbers and r > 0 else None
                if parsed:
                    cell.value, cell.number_format = parsed
                else:
                    cell.value = ILLEGAL_CHARACTERS_RE.sub("", value) or None
                    if len(value) > 60:  # textos longos quebram dentro da célula
                        cell.alignment = wrap
                if r == 0:
                    cell.font = bold
        return row + len(table) + 1  # uma linha em branco entre tabelas

    if one_sheet:
        ws = wb.create_sheet(tr("Tabelas"))
        row = 1
        for table in (t for tables in per_page for t in tables):
            row = write(ws, row, table)
    else:
        for i, tables in enumerate(per_page):
            if tables:
                ws = wb.create_sheet(tr("Página {0}", i + 1))
                row = 1
                for table in tables:
                    row = write(ws, row, table)

    for ws in wb.worksheets:  # largura das colunas pelo conteúdo
        widths: dict[int, int] = {}
        for cells in ws.iter_rows():
            for cell in cells:
                if cell.value is not None:
                    longest = max(len(part) for part in str(cell.value).split("\n"))
                    widths[cell.column] = max(widths.get(cell.column, 0), longest)
        for col, w in widths.items():
            ws.column_dimensions[get_column_letter(col)].width = min(max(w + 2, 8), 60)

    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    try:
        wb.save(out)
    except PermissionError as exc:
        raise PdfError(tr("Não foi possível salvar “{0}”.\nSe ele estiver aberto no Excel, feche e tente de novo.", os.path.basename(out))) from exc
    ctx.progress(1, 1, tr("Concluído"))
    return count
